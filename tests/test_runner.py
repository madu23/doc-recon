import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from docrecon.models import OCRResult
from docrecon.runner import index_images, run

ROOT=Path(__file__).parents[1]

def workbook(path):
    wb=Workbook(); ws=wb.active
    ws.append(["Synthetic DocRecon fixture"])
    ws.append(["S/N","FIRST_NAME","LAST_NAME","MOBILE","STREET","CITY","POD number"])
    ws.append([1,"Ada","Example","08031112233","12 Demo Avenue","Sample City","DEMO00000001"])
    wb.save(path)

def fake_extract(*args, **kwargs):
    return OCRResult(extracted_name="Ada Example", extracted_phone="08031112233",
        extracted_address="12 Demo Avenue", extracted_reference="DEMO00000001",
        receiver_name_present=True, signature_present=True, delivery_date_present=True,
        ocr_confidence=95)

def test_index_reports_duplicate(tmp_path):
    (tmp_path/"DEMO-1.jpg").write_bytes(b"x"); (tmp_path/"DEMO1.png").write_bytes(b"x")
    _, duplicates=index_images(tmp_path,[".jpg",".png"])
    assert duplicates == ["DEMO1"]

def test_duplicate_routes_to_review(tmp_path):
    source=tmp_path/"input.xlsx"; pods=tmp_path/"pods"; pods.mkdir(); workbook(source)
    (pods/"DEMO00000001.jpg").write_bytes(b"x"); (pods/"DEMO-00000001.png").write_bytes(b"x")
    cfg=json.loads((ROOT/"config.json").read_text())
    results,_=run(source,pods,tmp_path/"rejected",tmp_path/"output",cfg,dry_run=True)
    assert results[0].final_status == "REVIEW"
    assert results[0].reason_codes == ["DUPLICATE_REFERENCE"]

def test_dry_run_preserves_input_and_writes_report(tmp_path, monkeypatch):
    source=tmp_path/"input.xlsx"; pods=tmp_path/"pods"; pods.mkdir()
    image=pods/"DEMO00000001.jpg"; image.write_bytes(b"synthetic")
    workbook(source)
    monkeypatch.setattr("docrecon.runner.extract_pod", fake_extract)
    cfg=json.loads((ROOT/"config.json").read_text())
    results, report=run(source,pods,tmp_path/"rejected",tmp_path/"output",cfg,dry_run=True)
    assert image.exists() and report.exists()
    assert results[0].final_status == "SUCCESSFUL"
    assert "FULL RECORDS" in load_workbook(report, read_only=True).sheetnames

def test_copy_preserves_input(tmp_path, monkeypatch):
    source=tmp_path/"input.xlsx"; pods=tmp_path/"pods"; pods.mkdir()
    image=pods/"DEMO00000001.jpg"; image.write_bytes(b"synthetic")
    workbook(source); monkeypatch.setattr("docrecon.runner.extract_pod", fake_extract)
    cfg=json.loads((ROOT/"config.json").read_text())
    results,_=run(source,pods,tmp_path/"rejected",tmp_path/"output",cfg,copy_files=True)
    assert image.exists()
    assert (tmp_path/"output"/"accepted"/"1-DEMO00000001.jpg").exists()
