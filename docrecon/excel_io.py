from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from .models import SourceRecord, ValidationResult


def _string(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def read_source_records(path: Path, cfg: dict) -> list[SourceRecord]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    header_row = int(cfg["excel"].get("header_row", 1))
    headers = {str(cell.value).strip(): idx for idx, cell in enumerate(ws[header_row], start=1) if cell.value is not None}
    cols = cfg["excel"]["columns"]
    missing = [v for v in cols.values() if v not in headers]
    if missing:
        raise ValueError(f"Missing required Excel column(s): {', '.join(missing)}")
    records = []
    seen = set()
    for row_idx in range(header_row + 1, ws.max_row + 1):
        def val(key):
            return _string(ws.cell(row_idx, headers[cols[key]]).value)
        pod = val("pod_number")
        if not pod:
            continue
        rec = SourceRecord(
            row_number=row_idx,
            serial=val("serial"),
            first_name=val("first_name"),
            last_name=val("last_name"),
            phone=val("phone"),
            address=val("address"),
            location=val("location"),
            pod_number=pod,
        )
        if pod.upper() in seen:
            raise ValueError(f"Duplicate POD number in Excel: {pod}")
        seen.add(pod.upper())
        records.append(rec)
    return records


def write_report(results: list[ValidationResult], output_path: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "SUMMARY"

    total = len(results)
    successful = sum(1 for r in results if r.final_status == "SUCCESSFUL")
    review = sum(1 for r in results if r.final_status == "REVIEW")
    rejected = sum(1 for r in results if r.final_status == "REJECTED")
    not_found = sum(1 for r in results if "POD_NOT_FOUND" in r.reason_codes)
    mismatch = sum(1 for r in results if any(x in r.reason_codes for x in ["REFERENCE_MISMATCH", "IDENTITY_MISMATCH"]))
    metrics = [
        ("Metric", "Count"),
        ("Total records", total),
        ("Successful", successful),
        ("Review", review),
        ("Action Required", rejected + review),
        ("Rejected", rejected),
        ("POD Not Found", not_found),
        ("Customer/POD Mismatch", mismatch),
    ]
    ws["A1"] = "DocRecon — POD Validation Summary"
    ws["A1"].font = Font(size=16, bold=True)
    for i, (m, c) in enumerate(metrics, start=3):
        ws.cell(i, 1, m); ws.cell(i, 2, c)
    ws["A3"].font = ws["B3"].font = Font(bold=True)

    detail = wb.create_sheet("FULL RECORDS")
    headers = [
        "S/N", "FIRST_NAME", "LAST_NAME", "Customer", "Excel Original POD",
        "Final Verified POD", "POD Filename", "POD Serial Number", "POD Location",
        "Identity Status", "Completion Status", "Final Status", "Action Required",
        "Reason/Evidence", "OCR Name", "Name Score", "OCR Phone", "Phone Match",
        "OCR Address", "Address Score", "OCR Reference", "Reference Match",
        "Receiver Name Present", "Signature Present", "Delivery Date Present", "OCR Confidence",
        "Reason Codes", "OCR Engine", "OCR Engine Version", "OCR Time (ms)", "Total Time (ms)"
    ]
    detail.append(headers)
    for r in results:
        e = r.extracted
        detail.append([
            r.source.serial, r.source.first_name, r.source.last_name, r.source.customer_name,
            r.source.pod_number, e.extracted_reference or r.source.pod_number,
            r.final_filename or r.original_filename, r.source.serial if r.final_status == "SUCCESSFUL" else "",
            r.source.location, r.identity_status, r.completion_status, r.final_status,
            r.action_required, r.reason_evidence, e.extracted_name, round(r.name_score, 1),
            e.extracted_phone, "YES" if r.phone_match else "NO", e.extracted_address,
            round(r.address_score, 1), e.extracted_reference, "YES" if r.reference_match else "NO",
            "YES" if e.receiver_name_present else "NO", "YES" if e.signature_present else "NO",
            "YES" if e.delivery_date_present else "NO", round(e.ocr_confidence, 1),
            ", ".join(r.reason_codes), e.engine, e.engine_version,
            round(e.processing_time_ms, 1), round(r.processing_time_ms, 1)
        ])

    challenges = wb.create_sheet("CHALLENGES")
    challenges.append(["S/N", "Customer", "Problem", "What was found", "What is required", "Priority"])
    for r in results:
        if r.final_status == "SUCCESSFUL":
            continue
        problem = "; ".join(r.reason_codes) or "REVIEW REQUIRED"
        found = r.reason_evidence
        required = _required_action(r.reason_codes)
        priority = "HIGH" if any(x in r.reason_codes for x in ["REFERENCE_MISMATCH", "IDENTITY_MISMATCH", "POD_NOT_FOUND"]) else "MEDIUM"
        challenges.append([r.source.serial, r.source.customer_name, problem, found, required, priority])

    corrections = wb.create_sheet("CORRECTIONS")
    corrections.append(["S/N", "Customer", "Original POD", "Correct POD", "Reason", "Evidence", "Status"])
    # V1 deliberately does not auto-correct POD assignments. Any correction is a human-reviewed action.

    loc = wb.create_sheet("LOCATION SUMMARY")
    loc.append(["Location", "Successful PODs", "Exceptions"])
    locations = {}
    for r in results:
        key = r.source.location or "Unknown"
        d = locations.setdefault(key, [0, 0])
        if r.final_status == "SUCCESSFUL": d[0] += 1
        else: d[1] += 1
    for key in sorted(locations):
        loc.append([key, locations[key][0], locations[key][1]])

    # Styling and widths
    for sheet in wb.worksheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="D9EAD3")
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx in range(1, sheet.max_column + 1):
            max_len = 0
            for row_idx in range(1, min(sheet.max_row, 2000) + 1):
                v = sheet.cell(row_idx, col_idx).value
                if v is not None:
                    max_len = max(max_len, len(str(v)))
            sheet.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45)
        sheet.freeze_panes = "A2" if sheet.title != "SUMMARY" else "A3"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _required_action(codes: list[str]) -> str:
    if "POD_NOT_FOUND" in codes: return "Provide the corresponding POD image."
    if "REFERENCE_MISMATCH" in codes: return "Provide or verify the correct POD/reference."
    if "IDENTITY_MISMATCH" in codes: return "Verify customer identity against the correct POD."
    if any(x in codes for x in ["RECEIVER_NAME_MISSING", "SIGNATURE_MISSING", "DELIVERY_DATE_MISSING"]):
        return "Provide a sufficiently completed POD with receiver name, signature and delivery date."
    return "Manually review the POD and source data."
